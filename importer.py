"""Import people from a spreadsheet in the export's format.

Two steps: plan() reads the file and works out every add and change without
touching the database, so the page can show a preview; apply() then makes
exactly those changes in one transaction. Importing never deletes people.
"""
import re
from datetime import date, datetime

from openpyxl import load_workbook

from export import display_name

# Header text (lower-case) -> field. Columns not listed are ignored.
HEADERS = {
    'id': 'id', 'name': 'name', 'nickname': 'nickname', 'gender': 'gender',
    'birth date': 'birth_date', 'death date': 'death_date', 'status': 'status',
    'birth order': 'birth_order', 'father': 'father', 'mother': 'mother',
    'spouses': 'spouses', 'spouse': 'spouses', 'notes': 'notes',
}
MAX_ROWS = 5000

GENDERS = {'male': 'Male', 'm': 'Male', 'lelaki': 'Male', 'female': 'Female', 'f': 'Female',
           'perempuan': 'Female', 'other': 'Other'}
DECEASED = {'meninggal dunia', 'meninggal', 'deceased', 'dead', 'arwah', 'allahyarham',
            'allahyarhamah', 'yes', 'y', '1', 'true'}
LIVING = {'', 'living', 'alive', 'hidup', 'no', 'n', '0', 'false'}

FIELD_LABELS = {
    'name': 'Name', 'nickname': 'Nickname', 'gender': 'Gender', 'birth_date': 'Birth date',
    'death_date': 'Death date', 'is_deceased': 'Meninggal Dunia', 'birth_order': 'Birth order',
    'father_id': 'Father', 'mother_id': 'Mother', 'notes': 'Notes',
}


class RowError(Exception):
    pass


def _text(value):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    value = str(value).strip()
    return value or None


def _date(value):
    if value is None or value == '':
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    s = str(value).strip()
    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%d.%m.%Y'):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            pass
    raise RowError(f'"{s}" is not a date (use YYYY-MM-DD or DD/MM/YYYY)')


def _positive_int(value, label):
    s = _text(value)
    if s is None:
        return None
    if not s.isdigit() or int(s) < 1:
        raise RowError(f'{label} must be a whole number from 1 up, not "{s}"')
    return int(s)


def _key(name):
    return re.sub(r'\s+', ' ', display_name(name or '')).lower()


def read_rows(file):
    wb = load_workbook(file, read_only=True, data_only=True)
    ws = wb['Family'] if 'Family' in wb.sheetnames else wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None) or ()
    columns = {}
    for i, h in enumerate(header):
        field = HEADERS.get((_text(h) or '').lower())
        if field and field not in columns:
            columns[field] = i
    if 'name' not in columns:
        raise RowError('The first row must be the column headers, including "Name"')

    out = []
    for n, row in enumerate(rows, start=2):
        if n - 1 > MAX_ROWS:
            raise RowError(f'Too many rows (the limit is {MAX_ROWS})')
        cells = {f: (row[i] if i < len(row) else None) for f, i in columns.items()}
        if all(_text(v) is None for v in cells.values()):
            continue
        out.append((n, cells))
    return columns, out


def plan(file, conn):
    try:
        columns, rows = read_rows(file)
    except RowError as e:
        return {'errors': [{'row': 1, 'message': str(e)}]}
    except Exception:
        return {'errors': [{'row': 0, 'message': 'Could not read that file. Save it as .xlsx and try again.'}]}

    members = {r['id']: dict(r) for r in conn.execute('SELECT * FROM members')}
    marriages = {(r['person_a'], r['person_b']) for r in conn.execute('SELECT person_a, person_b FROM marriages')}
    errors = []

    # 1. Each row becomes a target: an existing id, or ('new', row number).
    parsed = []
    seen_ids = set()
    for n, cells in rows:
        try:
            target = None
            if 'id' in cells and _text(cells['id']) is not None:
                s = _text(cells['id']).lstrip('#')
                if not s.isdigit() or int(s) not in members:
                    raise RowError(f'ID {s} is not anyone in the tree (leave ID blank to add someone new)')
                target = int(s)
                if target in seen_ids:
                    raise RowError(f'ID {target} appears on more than one row')
                seen_ids.add(target)
            name = _text(cells.get('name'))
            if not name:
                raise RowError('Name is required')
            parsed.append({'row': n, 'cells': cells, 'target': target if target else ('new', n), 'name': name})
        except RowError as e:
            errors.append({'row': n, 'message': str(e)})

    # 2. Who a name refers to: the sheet's names, then the tree's names for
    #    everyone not on the sheet.
    names = {}
    for p in parsed:
        names.setdefault(_key(p['name']), set()).add(p['target'])
    for i, m in members.items():
        if i not in seen_ids:
            names.setdefault(_key(m['name']), set()).add(i)

    def resolve(value, what):
        s = _text(value)
        if s is None:
            return None
        m = re.match(r'^#(\d+)\b', s)
        if m:
            i = int(m.group(1))
            if i not in members:
                raise RowError(f'{what} "{s}": there is no #{i} in the tree')
            return i
        found = names.get(_key(s), set())
        if not found:
            raise RowError(f'{what} "{s}" is not on the sheet or in the tree')
        if len(found) > 1:
            raise RowError(f'{what} "{s}" matches more than one person; write #ID instead')
        return next(iter(found))

    # 3. Field values per row.
    for p in parsed:
        c = p['cells']
        try:
            v = {}
            v['name'] = p['name']
            if 'nickname' in c: v['nickname'] = _text(c['nickname'])
            if 'notes' in c: v['notes'] = _text(c['notes'])
            if 'gender' in c:
                g = _text(c['gender'])
                if g is not None and g.lower() not in GENDERS:
                    raise RowError(f'Gender "{g}" should be Male, Female or Other')
                v['gender'] = GENDERS[g.lower()] if g else None
            if 'birth_date' in c: v['birth_date'] = _date(c['birth_date'])
            if 'death_date' in c: v['death_date'] = _date(c['death_date'])
            if 'birth_order' in c: v['birth_order'] = _positive_int(c['birth_order'], 'Birth order')
            if 'status' in c:
                st = (_text(c['status']) or '').lower()
                if st not in DECEASED and st not in LIVING:
                    raise RowError(f'Status "{st}" should be "Meninggal Dunia" or blank')
                v['is_deceased'] = 1 if st in DECEASED else 0
            if v.get('death_date'):
                v['is_deceased'] = 1
            if v.get('birth_date') and v.get('death_date') and v['death_date'] < v['birth_date']:
                raise RowError('Death date is before birth date')
            for field, col in (('father_id', 'father'), ('mother_id', 'mother')):
                if col in c:
                    ref = resolve(c[col], col.capitalize())
                    if ref == p['target']:
                        raise RowError(f'Someone cannot be their own {col}')
                    v[field] = ref
            if 'spouses' in c:
                refs = []
                for part in re.split(r'[,;]', _text(c['spouses']) or ''):
                    if part.strip():
                        ref = resolve(part, 'Spouse')
                        if ref == p['target']:
                            raise RowError('Someone cannot be their own spouse')
                        refs.append(ref)
                p['spouses'] = refs
            p['values'] = v
        except RowError as e:
            errors.append({'row': p['row'], 'message': str(e)})

    if errors:
        return {'errors': sorted(errors, key=lambda e: e['row'])}

    # 4. What would change.
    label = lambda t: p_by_target[t]['name'] if t in p_by_target else display_name(members[t]['name'])
    p_by_target = {p['target']: p for p in parsed}
    adds, updates = [], []
    for p in parsed:
        if isinstance(p['target'], tuple):
            adds.append({'row': p['row'], 'name': display_name(p['name'])})
            continue
        old = members[p['target']]
        changes = []
        for field, new in p['values'].items():
            before = old.get(field)
            # The export writes names in display form ("Ahmad Faiz", not
            # "ahmad faiz"); that alone is not a change worth making.
            if field == 'name' and _key(before) == _key(new):
                p['values'][field] = before
                continue
            if (before or None) != (new or None):
                show = lambda x: label(x) if field in ('father_id', 'mother_id') and x else x
                changes.append({'field': FIELD_LABELS.get(field, field),
                                'from': show(before), 'to': show(new)})
        if changes:
            updates.append({'row': p['row'], 'name': display_name(p['name']), 'changes': changes})

    pair = lambda a, b: (a, b) if str(a) < str(b) else (b, a)
    wanted = {pair(p['target'], s) for p in parsed for s in p.get('spouses', [])}
    existing = {pair(a, b) for a, b in marriages}
    sheet_people = {p['target'] for p in parsed if 'spouses' in p}
    marriages_added = [w for w in wanted if w not in existing]
    # A marriage is only removed when both people are on the sheet and neither lists the other.
    marriages_removed = [e for e in existing
                         if e[0] in sheet_people and e[1] in sheet_people and e not in wanted]

    unchanged = len(parsed) - len(adds) - len(updates)
    return {
        'errors': [],
        'adds': adds,
        'updates': updates,
        'marriages_added': [[label(a), label(b)] for a, b in marriages_added],
        'marriages_removed': [[label(a), label(b)] for a, b in marriages_removed],
        'unchanged': unchanged,
        '_parsed': parsed,
        '_marriages': (marriages_added, marriages_removed),
    }


def apply(result, conn):
    """Make the changes plan() described, in one transaction."""
    parsed = result['_parsed']
    new_ids = {}
    for p in parsed:
        if isinstance(p['target'], tuple):
            cur = conn.execute('INSERT INTO members (name) VALUES (?)', (p['name'],))
            new_ids[p['target']] = cur.lastrowid
    real = lambda t: new_ids.get(t, t) if isinstance(t, tuple) else t

    for p in parsed:
        values = {k: (real(v) if k in ('father_id', 'mother_id') and v is not None else v)
                  for k, v in p['values'].items()}
        cols = ', '.join(f'{k}=?' for k in values)
        conn.execute(f'UPDATE members SET {cols} WHERE id=?', [*values.values(), real(p['target'])])

    added, removed = result['_marriages']
    for a, b in added:
        a, b = sorted((real(a), real(b)))
        conn.execute('INSERT OR IGNORE INTO marriages (person_a, person_b) VALUES (?, ?)', (a, b))
    for a, b in removed:
        a, b = sorted((real(a), real(b)))
        conn.execute('DELETE FROM marriages WHERE person_a=? AND person_b=?', (a, b))
    return len(new_ids)
