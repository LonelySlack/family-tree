"""Spreadsheet export of the whole tree, in the format importer.py reads back."""
import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

LOWERCASE_WORDS = {'bin', 'binti', 'bt', 'bte', 'b.', 'al', 'van', 'de'}


def display_name(name):
    """Same formatting as the web page: "ahmad faiz bin abu bakar" -> "Ahmad Faiz bin Abu Bakar"."""
    words = (name or '').split()
    return ' '.join(
        w.lower() if i > 0 and w.lower() in LOWERCASE_WORDS else w[:1].upper() + w[1:].lower()
        for i, w in enumerate(words)
    )


def age(member, today=None):
    if not member['birth_date']:
        return None
    if member['is_deceased'] and not member['death_date']:
        return None
    born = date.fromisoformat(member['birth_date'])
    end = date.fromisoformat(member['death_date']) if member['death_date'] else (today or date.today())
    return end.year - born.year - ((end.month, end.day) < (born.month, born.day))


def generations(members):
    """Generation number per id, 1 = oldest, matching the tree's rows:
    a child sits below both parents, and spouses share a generation."""
    level = {m['id']: 0 for m in members}
    for _ in range(100):
        changed = False
        for m in members:
            new = level[m['id']]
            for p in (m['father_id'], m['mother_id']):
                if p in level:
                    new = max(new, level[p] + 1)
            for s in m.get('spouse_ids', []):
                if s in level:
                    new = max(new, level[s])
            if new != level[m['id']]:
                level[m['id']] = new
                changed = True
        if not changed:
            break
    return {i: lv + 1 for i, lv in level.items()}


# Columns the importer reads are marked; the rest are for reading only.
COLUMNS = [
    ('ID', 6),            # read: which existing person a row updates
    ('Generation', 11),
    ('Name', 34),         # read
    ('Nickname', 16),     # read
    ('Gender', 9),        # read
    ('Birth date', 12),   # read
    ('Death date', 12),   # read
    ('Status', 16),       # read
    ('Birth order', 11),  # read
    ('Age', 6),
    ('Father', 30),       # read
    ('Mother', 30),       # read
    ('Spouses', 34),      # read
    ('Children', 50),
    ('Notes', 50),        # read
]
WRAP = {'Spouses', 'Children', 'Notes'}


def _style(ws):
    header_fill = PatternFill('solid', start_color='0D6EFD')
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = header_fill
        cell.alignment = Alignment(vertical='center')
    names = [c for c, _ in COLUMNS]
    for i, (_, width) in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=names[cell.column - 1] in WRAP)
    ws.freeze_panes = 'D2'
    ws.auto_filter.ref = ws.dimensions


def _instructions(wb):
    ws = wb.create_sheet('How to import')
    lines = [
        'How to fill this in',
        '',
        'One row per person, on the "Family" sheet. Only the first sheet is read.',
        'ID: leave blank for someone new. Keep the ID to update someone already in the tree.',
        'Name is required. Gender: Male, Female or Other.',
        'Dates: YYYY-MM-DD (e.g. 1962-08-30) or DD/MM/YYYY.',
        'Status: "Meninggal Dunia" or "Deceased" if they have passed away; otherwise leave blank.',
        'Birth order: 1 for the eldest child, 2 for the next, and so on.',
        'Father / Mother: their name exactly as written in the Name column, or #ID (e.g. #12).',
        'Spouses: names or #IDs separated by commas.',
        'Generation, Age and Children are worked out by the app and ignored when importing.',
        'Importing never deletes anyone; you will see a preview before anything is saved.',
    ]
    for line in lines:
        ws.append([line])
    ws['A1'].font = Font(bold=True, size=13)
    ws.column_dimensions['A'].width = 100


def to_xlsx(members):
    by_id = {m['id']: m for m in members}
    gen = generations(members)
    name_count = {}
    for m in members:
        key = display_name(m['name']).lower()
        name_count[key] = name_count.get(key, 0) + 1

    def ref(i):
        """A person as the importer will read them back: their name, or
        "#ID name" when two people share that name."""
        if i not in by_id:
            return ''
        name = display_name(by_id[i]['name'])
        return name if name_count[name.lower()] == 1 else f'#{i} {name}'

    wb = Workbook()
    ws = wb.active
    ws.title = 'Family'
    ws.append([c for c, _ in COLUMNS])

    ordered = sorted(members, key=lambda m: (
        gen[m['id']], m['birth_order'] or 999, m['birth_date'] or '9999', m['id']))
    for m in ordered:
        children = [c for c in ordered if m['id'] in (c['father_id'], c['mother_id'])]
        ws.append([
            m['id'],
            gen[m['id']],
            display_name(m['name']),
            m['nickname'] or '',
            m['gender'] or '',
            m['birth_date'] or '',
            m['death_date'] or '',
            'Meninggal Dunia' if m['is_deceased'] else '',
            m['birth_order'],
            age(m),
            ref(m['father_id']),
            ref(m['mother_id']),
            ', '.join(ref(s) for s in m.get('spouse_ids', [])),
            ', '.join(display_name(c['name']) for c in children),
            m['notes'] or '',
        ])
    _style(ws)
    _instructions(wb)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def template_xlsx():
    """Empty sheet with the right headers, for starting a tree from Excel."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Family'
    ws.append([c for c, _ in COLUMNS])
    _style(ws)
    _instructions(wb)
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out
